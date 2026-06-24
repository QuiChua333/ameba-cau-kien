"""
excel_service.py — Fill the 掘削深度_Templete sheet in the Excel template
with extracted foundation/beam data and return the file as bytes.
"""
import io
import re
from copy import copy
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.datavalidation import DataValidation

from models import FoundationItem, PitHoleItem

# The "Loại Móng" (column D) dropdown in the template is an Excel x14-extension list
# validation that openpyxl silently DROPS on load (→ exported file loses the dropdown).
# We re-add it as a standard list validation pointing at the same source range on the
# データ sheet so the dropdown survives the export.
_LOAI_MONG_LIST = "データ!$A$2:$A$12"
_LOAI_MONG_COL = "D"

TEMPLATE_PATH = Path(__file__).parent.parent / "templates" / "計算書(施工) - DD.xlsx"

# Note: a floor-plan vs section 天端 conflict (FoundationItem.top_elevation_alt) is
# surfaced in the web UI only. The Excel export stays plain — no fill, no note.

# Mapping: Gemini classification → Excel "Loại Móng" string
CLASSIFICATION_MAP = {
    "DD":    "TNF-DD",
    "D":     "TNF-D",
    "TNF":   "TNF",
    "FW/FG": "FW/FG",
}

# The data table starts at this row (header is row 9)
DATA_START_ROW = 10

# Last column to replicate when extending the table (A–P = 1–16)
_TABLE_MAX_COL = 16


def _last_formatted_row(ws, start: int, col: int = 5, hard_cap: int = 2000) -> int:
    """Return the last pre-formatted data row in the template.

    Template data rows carry a formula in column E (VLOOKUP). The formatted block
    is the contiguous run of such rows starting at `start`; the first gap ends it.
    Falls back to `start` when the template has no formula rows.
    """
    last = start - 1
    for r in range(start, hard_cap):
        c = ws.cell(row=r, column=col)
        if isinstance(c.value, str) and c.value.startswith("="):
            last = r
        elif last >= start:
            break
    return max(last, start)


def _clone_row_format(ws, src_row: int, dst_row: int, max_col: int = _TABLE_MAX_COL):
    """Replicate a template row's styling and formulas onto a new (overflow) row.

    Copies the full cell style for every column so borders/fills/fonts match the
    template, and re-creates the computed columns (E, J, K, L) by translating each
    formula's relative references from the source row to the destination row.
    Literal data-column values are NOT copied — fill_excel writes those itself.
    """
    for col in range(1, max_col + 1):
        s = ws.cell(row=src_row, column=col)
        d = ws.cell(row=dst_row, column=col)
        if s.has_style:
            d._style = copy(s._style)
        if isinstance(s.value, str) and s.value.startswith("="):
            d.value = Translator(s.value, origin=s.coordinate).translate_formula(d.coordinate)
    # Match row height so extended rows look identical to the template.
    src_h = ws.row_dimensions[src_row].height
    if src_h is not None:
        ws.row_dimensions[dst_row].height = src_h


def _parse_d(d_value: str) -> tuple[float | None, float | None]:
    """
    Parse D dimension string into (d_main, d2).

    Supported formats:
      "700~300"     → (700, 300)
      "900(~250)"   → (900, 250)   ← parenthetical format from PDF/pdfplumber
      "900（~250）"  → (900, 250)   ← full-width parentheses variant
      "400"         → (400, None)
      "0"           → (0,   None)
    """
    def _to_float(t: str) -> float:
        return float(t.strip().replace(',', '').replace('，', ''))

    s = str(d_value).strip()

    # Format: "900(~250)" or "900（~250）" — parenthetical range
    m = re.search(r'^([\d,，]+)\s*[（(]~?\s*([\d,，]+)[)）]', s)
    if m:
        try:
            return _to_float(m.group(1)), _to_float(m.group(2))
        except ValueError:
            pass

    # Format: "700~300" — tilde-separated range
    if "~" in s:
        parts = s.split("~", 1)
        try:
            return _to_float(parts[0]), _to_float(parts[1])
        except ValueError:
            pass

    # Plain number
    try:
        return _to_float(s), None
    except ValueError:
        return None, None


def fill_excel(
    foundation_list: List[FoundationItem],
    pit_list: List[PitHoleItem] | None = None,
) -> bytes:
    """
    Load the Excel template, fill the 掘削深度_Templete sheet with
    foundation, beam data, and readable pit holes, then return as bytes.
    """
    wb = openpyxl.load_workbook(str(TEMPLATE_PATH))
    ws = wb["掘削深度_Templete"]

    # Last row the template pre-formats (borders + E/J/K/L formulas). Rows beyond
    # this are cloned from it so a long list never spills into unformatted cells.
    last_fmt_row = _last_formatted_row(ws, DATA_START_ROW)

    def _ensure_formatted(row: int):
        if row > last_fmt_row:
            _clone_row_format(ws, last_fmt_row, row)

    # Filter: only foundations and beams (not empty rows)
    items = [item for item in foundation_list if item.type]

    for i, item in enumerate(items):
        row = DATA_START_ROW + i
        _ensure_formatted(row)

        # A — Zone (default 1)
        ws.cell(row=row, column=1).value = 1

        # B — STT (sequential number)
        ws.cell(row=row, column=2).value = i + 1

        # C — Tên Móng (foundation/beam name)
        ws.cell(row=row, column=3).value = item.type

        # D — Loại Móng (mapped classification)
        loai_mong = CLASSIFICATION_MAP.get(item.classification, item.classification)
        ws.cell(row=row, column=4).value = loai_mong

        # F — Cao độ mặt trên cấu kiện H1 (top_elevation, mm)
        # G — H2 (for DD type, same as H1 unless we have specific data)
        # The 伏図 value is used as-is. A floor-plan/section conflict (top_elevation_alt)
        # is flagged in the web UI only — the Excel file stays plain (no fill, no note).
        if item.top_elevation is not None:
            ws.cell(row=row, column=6, value=item.top_elevation)
            ws.cell(row=row, column=7, value=item.top_elevation)

        # H — Chiều cao cấu kiện D (main height)
        # I — Chiều cao D2 Daike (only for D / DD)
        d_main, d2 = _parse_d(item.dimensions.D)
        if d_main is not None and d_main > 0:
            ws.cell(row=row, column=8).value = d_main   # H
        if d2 is not None and item.classification in ("D", "DD"):
            ws.cell(row=row, column=9).value = d2        # I

    # ── Readable pit holes (appended after foundation rows) ──────────────────
    readable_pits = [p for p in (pit_list or []) if p.readable and p.type]
    pit_offset = len(items)
    for j, pit in enumerate(readable_pits):
        row = DATA_START_ROW + pit_offset + j
        _ensure_formatted(row)

        ws.cell(row=row, column=1).value = 1                      # A — Zone
        ws.cell(row=row, column=2).value = pit_offset + j + 1     # B — STT
        ws.cell(row=row, column=3).value = pit.type               # C — Name
        ws.cell(row=row, column=4).value = "ピット"               # D — Type

        if pit.top_elevation is not None:
            ws.cell(row=row, column=6).value = pit.top_elevation  # F — H1
            ws.cell(row=row, column=7).value = pit.top_elevation  # G — H2

        if pit.D is not None and pit.D > 0:
            ws.cell(row=row, column=8).value = pit.D              # H — D

    # Restore the "Loại Móng" dropdown openpyxl dropped on load, covering every data
    # row that was written (including cloned overflow rows).
    last_data_row = DATA_START_ROW + len(items) + len(readable_pits) - 1
    if last_data_row >= DATA_START_ROW:
        dv = DataValidation(type="list", formula1=_LOAI_MONG_LIST, allowBlank=True)
        dv.add(f"{_LOAI_MONG_COL}{DATA_START_ROW}:{_LOAI_MONG_COL}{last_data_row}")
        ws.add_data_validation(dv)

    # Save to bytes buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
