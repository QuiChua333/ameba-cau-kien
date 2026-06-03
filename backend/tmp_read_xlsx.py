import openpyxl

wb = openpyxl.load_workbook(
    r'c:/quihn/ameba1/source/backend/templates/計算書(施工) - DD.xlsx',
    data_only=False
)

print("=== Sheet: データ ===")
ws = wb['データ']
for row in ws.iter_rows(min_row=1, max_row=15):
    for cell in row:
        if cell.value is not None:
            print(f"  {cell.coordinate}: {repr(cell.value)}")

print()
print("=== Sheet: 掘削深度_Templete - Row 9 (headers) ===")
ws2 = wb['掘削深度_Templete']
for cell in ws2[9]:
    if cell.value:
        from openpyxl.utils import get_column_letter
        print(f"  Col {get_column_letter(cell.column)} ({cell.coordinate}): {cell.value}")
